"""
Genera index.html (dashboard estático) a partir de Detalle_Casos_MDI_2026.xlsx
Se corre DESPUÉS de exportar_casos_jira.py, en el mismo flujo de GitHub Actions.
"""

import json
from pathlib import Path
from datetime import datetime
import openpyxl

CARPETA = Path(__file__).resolve().parent
RUTA_EXCEL = CARPETA / "Detalle_Casos_MDI_2026.xlsx"
RUTA_HTML = CARPETA / "index.html"


def cargar_datos():
    wb = openpyxl.load_workbook(RUTA_EXCEL, data_only=True)
    ws = wb["Detalle de Casos"]
    encabezados = [c.value for c in ws[1]]
    filas = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0]:
            filas.append(dict(zip(encabezados, row)))
    return filas


def normalizar(texto):
    return str(texto or "").strip().lower()


def buscar_campo(fila, *candidatos):
    """
    Busca el valor de un campo en la fila probando varios nombres posibles,
    sin importar mayúsculas/minúsculas o espacios (ej. 'Status' vs 'status').
    """
    claves_normalizadas = {normalizar(k): k for k in fila.keys()}
    for candidato in candidatos:
        clave_real = claves_normalizadas.get(normalizar(candidato))
        if clave_real is not None:
            return fila.get(clave_real)
    return None


def contar_por(filas, *candidatos):
    conteo = {}
    for f in filas:
        valor = buscar_campo(f, *candidatos) or "(Sin dato)"
        conteo[valor] = conteo.get(valor, 0) + 1
    return dict(sorted(conteo.items(), key=lambda x: -x[1]))


def es_abierto_fila(fila):
    """
    Determina si Jira considera el ticket abierto.

    Prioridad:
      1) Status Category Key / Status Category (si el Excel los contiene).
      2) Resolution/Resolved como respaldo.
      3) Nombre del Status como compatibilidad con el Excel actual.

    Así no dependemos únicamente de una lista rígida de nombres de estados.
    """
    categoria_key = buscar_campo(
        fila,
        "Status Category Key", "Status category key",
        "StatusCategoryKey", "Categoría de estado key"
    )
    if categoria_key:
        ck = normalizar(categoria_key)
        if ck in ("done", "complete", "completed"):
            return False
        if ck in ("new", "indeterminate", "in progress", "in_progress", "todo", "to do"):
            return True

    categoria = buscar_campo(
        fila,
        "Status Category", "Status category",
        "Categoría de estado", "Categoria de estado"
    )
    if categoria:
        c = normalizar(categoria)
        if c in ("done", "listo", "finalizado", "cerrado", "complete", "completed"):
            return False
        if c in ("to do", "por hacer", "in progress", "en curso", "new", "indeterminate"):
            return True

    resolucion = buscar_campo(fila, "Resolution", "Resolución", "Resolucion")
    fecha_resolucion = buscar_campo(
        fila, "Resolved", "Resolution Date", "Fecha Resolución", "Fecha Resolucion"
    )
    if resolucion or fecha_resolucion:
        return False

    estado = buscar_campo(fila, "Status", "Estado")
    if not estado:
        return False

    e = normalizar(estado)
    return e not in (
        "finalizado", "cerrado", "no aplica", "resuelto",
        "done", "closed", "resolved", "cancelado", "cancelled"
    )


def armar_dataset(filas):
    total = len(filas)
    abiertos = [f for f in filas if es_abierto_fila(f)]
    cerrados = total - len(abiertos)

    por_estado = contar_por(filas, "Status", "Estado")
    por_tipo = contar_por(filas, "Issue Type", "Issue type", "Tipo")
    por_distribuidor = contar_por(filas, "Distribuidor")
    por_asignado_total = contar_por(filas, "Assignee", "Asignado")
    por_asignado_abiertos = contar_por(abiertos, "Assignee", "Asignado")

    # Casos abiertos más antiguos (por fecha de creación)
    def fecha_creacion(f):
        v = buscar_campo(f, "Created", "Creado", "Fecha Creación")
        if isinstance(v, datetime):
            return v
        try:
            return datetime.strptime(str(v)[:10], "%Y-%m-%d")
        except Exception:
            return datetime.now()

    abiertos_ordenados = sorted(abiertos, key=fecha_creacion)
    hoy = datetime.now()
    casos_antiguos = []
    for f in abiertos_ordenados[:30]:
        fc = fecha_creacion(f)
        dias = (hoy - fc).days
        casos_antiguos.append({
            "ticket": buscar_campo(f, "Ticket"),
            "tipo": buscar_campo(f, "Issue Type", "Issue type", "Tipo"),
            "distribuidor": buscar_campo(f, "Distribuidor") or "",
            "estado": buscar_campo(f, "Status", "Estado"),
            "asignado": buscar_campo(f, "Assignee", "Asignado") or "(Sin asignar)",
            "dias_abierto": dias,
            "creado": str(buscar_campo(f, "Created", "Creado", "Fecha Creación"))[:10],
        })
    casos_antiguos = sorted(casos_antiguos, key=lambda x: -x["dias_abierto"])

    detalle = []
    for f in filas:
        detalle.append({
            "ticket": buscar_campo(f, "Ticket"),
            "tipo": buscar_campo(f, "Issue Type", "Issue type", "Tipo"),
            "distribuidor": buscar_campo(f, "Distribuidor") or "",
            "estado": buscar_campo(f, "Status", "Estado"),
            "asignado": buscar_campo(f, "Assignee", "Asignado") or "",
            "creado": str(buscar_campo(f, "Created", "Creado", "Fecha Creación"))[:10],
            "resuelto": (str(buscar_campo(f, "Resolved", "Resolution Date", "Fecha Resolución"))[:10]
                         if buscar_campo(f, "Resolved", "Resolution Date", "Fecha Resolución") else ""),
        })

    return {
        "generado": hoy.strftime("%Y-%m-%d %H:%M"),
        "overview": {
            "total": total,
            "abiertos": len(abiertos),
            "cerrados": cerrados,
        },
        "por_estado": por_estado,
        "por_tipo": por_tipo,
        "por_distribuidor": por_distribuidor,
        "por_asignado_total": por_asignado_total,
        "por_asignado_abiertos": por_asignado_abiertos,
        "casos_antiguos": casos_antiguos,
        "detalle": detalle,
    }


HTML_TEMPLATE = """<!DOCTYPE html>
<html lang="es">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>Radar Postventa ISOFT</title>
<script src="https://cdnjs.cloudflare.com/ajax/libs/Chart.js/4.4.0/chart.umd.min.js"></script>
<style>
  * { box-sizing: border-box; }
  body { font-family: -apple-system, Segoe UI, Roboto, Arial, sans-serif; background: #f4f5f7; margin: 0; color: #172b4d; }
  header { background: #0747a6; color: white; padding: 20px 30px; }
  header h1 { margin: 0; font-size: 22px; }
  header p { margin: 4px 0 0; opacity: 0.85; font-size: 13px; }
  .container { padding: 24px; max-width: 1400px; margin: 0 auto; }
  .kpis { display: grid; grid-template-columns: repeat(auto-fit, minmax(180px, 1fr)); gap: 16px; margin-bottom: 24px; }
  .kpi { background: white; border-radius: 10px; padding: 18px; box-shadow: 0 1px 3px rgba(0,0,0,0.08); }
  .kpi .num { font-size: 30px; font-weight: 700; }
  .kpi .label { font-size: 13px; color: #6b778c; margin-top: 4px; }
  .grid { display: grid; grid-template-columns: repeat(auto-fit, minmax(380px, 1fr)); gap: 20px; margin-bottom: 24px; }
  .card { background: white; border-radius: 10px; padding: 18px; box-shadow: 0 1px 3px rgba(0,0,0,0.08); }
  .card h3 { margin: 0 0 14px; font-size: 15px; color: #172b4d; }
  table { width: 100%; border-collapse: collapse; font-size: 13px; }
  th, td { text-align: left; padding: 8px 10px; border-bottom: 1px solid #eee; }
  th { color: #6b778c; font-weight: 600; position: sticky; top: 0; background: white; }
  .badge { padding: 2px 8px; border-radius: 10px; font-size: 11px; background: #ffebe6; color: #bf2600; }
  input[type=text] { width: 100%; padding: 8px 10px; border: 1px solid #ddd; border-radius: 6px; margin-bottom: 10px; font-size: 13px; }
  .scroll { max-height: 420px; overflow-y: auto; }
</style>
</head>
<body>
<header>
  <h1>Radar Postventa ISOFT</h1>
  <p>Actualizado automáticamente el __GENERADO__ (proyecto MDI, casos 2026)</p>
</header>
<div class="container">

  <div class="kpis">
    <div class="kpi"><div class="num">__TOTAL__</div><div class="label">Casos totales 2026</div></div>
    <div class="kpi"><div class="num">__ABIERTOS__</div><div class="label">Casos abiertos</div></div>
    <div class="kpi"><div class="num">__CERRADOS__</div><div class="label">Casos cerrados</div></div>
  </div>

  <div class="grid">
    <div class="card"><h3>Casos por estado</h3><canvas id="chartEstado"></canvas></div>
    <div class="card"><h3>Casos por tipo</h3><canvas id="chartTipo"></canvas></div>
    <div class="card"><h3>Casos por distribuidor (top 10)</h3><canvas id="chartDistribuidor"></canvas></div>
    <div class="card"><h3>Carga por asignado - abiertos (top 10)</h3><canvas id="chartAsignado"></canvas></div>
  </div>

  <div class="card" style="margin-bottom:24px;">
    <h3>Casos abiertos más antiguos (top 30)</h3>
    <div class="scroll">
    <table>
      <thead><tr><th>Ticket</th><th>Tipo</th><th>Distribuidor</th><th>Estado</th><th>Asignado</th><th>Días abierto</th><th>Creado</th></tr></thead>
      <tbody id="tablaAntiguos"></tbody>
    </table>
    </div>
  </div>

  <div class="card">
    <h3>Detalle completo (filtrable)</h3>
    <input type="text" id="filtroDetalle" placeholder="Buscar por ticket, distribuidor, estado, asignado...">
    <div class="scroll">
    <table>
      <thead><tr><th>Ticket</th><th>Tipo</th><th>Distribuidor</th><th>Estado</th><th>Asignado</th><th>Creado</th><th>Resuelto</th></tr></thead>
      <tbody id="tablaDetalle"></tbody>
    </table>
    </div>
  </div>

</div>

<script>
const DATA = __DATA_JSON__;

function topN(obj, n) {
  return Object.entries(obj).slice(0, n);
}

function hacerBarra(id, labels, valores, color) {
  new Chart(document.getElementById(id), {
    type: 'bar',
    data: { labels: labels, datasets: [{ data: valores, backgroundColor: color }] },
    options: { indexAxis: 'y', plugins: { legend: { display: false } }, responsive: true }
  });
}

hacerBarra('chartEstado', Object.keys(DATA.por_estado), Object.values(DATA.por_estado), '#0052cc');
hacerBarra('chartTipo', Object.keys(DATA.por_tipo), Object.values(DATA.por_tipo), '#00875a');

const distTop = topN(DATA.por_distribuidor, 10);
hacerBarra('chartDistribuidor', distTop.map(x=>x[0]), distTop.map(x=>x[1]), '#ff8b00');

const asigTop = topN(DATA.por_asignado_abiertos, 10);
hacerBarra('chartAsignado', asigTop.map(x=>x[0]), asigTop.map(x=>x[1]), '#de350b');

const tbodyAntiguos = document.getElementById('tablaAntiguos');
DATA.casos_antiguos.forEach(c => {
  const tr = document.createElement('tr');
  tr.innerHTML = `<td>${c.ticket}</td><td>${c.tipo||''}</td><td>${c.distribuidor||''}</td><td>${c.estado||''}</td><td>${c.asignado||''}</td><td><span class="badge">${c.dias_abierto} días</span></td><td>${c.creado||''}</td>`;
  tbodyAntiguos.appendChild(tr);
});

const tbodyDetalle = document.getElementById('tablaDetalle');
function pintarDetalle(filtro) {
  tbodyDetalle.innerHTML = '';
  const f = (filtro || '').toLowerCase();
  DATA.detalle
    .filter(c => !f || Object.values(c).join(' ').toLowerCase().includes(f))
    .slice(0, 500)
    .forEach(c => {
      const tr = document.createElement('tr');
      tr.innerHTML = `<td>${c.ticket}</td><td>${c.tipo||''}</td><td>${c.distribuidor||''}</td><td>${c.estado||''}</td><td>${c.asignado||''}</td><td>${c.creado||''}</td><td>${c.resuelto||''}</td>`;
      tbodyDetalle.appendChild(tr);
    });
}
pintarDetalle('');
document.getElementById('filtroDetalle').addEventListener('input', e => pintarDetalle(e.target.value));
</script>
</body>
</html>
"""


def generar_html():
    filas = cargar_datos()
    if filas:
        print(f"Columnas detectadas en el Excel: {list(filas[0].keys())}")
    dataset = armar_dataset(filas)
    print(f"Total: {dataset['overview']['total']}  Abiertos: {dataset['overview']['abiertos']}  Cerrados: {dataset['overview']['cerrados']}")

    html = HTML_TEMPLATE
    html = html.replace("__GENERADO__", dataset["generado"])
    html = html.replace("__TOTAL__", str(dataset["overview"]["total"]))
    html = html.replace("__ABIERTOS__", str(dataset["overview"]["abiertos"]))
    html = html.replace("__CERRADOS__", str(dataset["overview"]["cerrados"]))
    html = html.replace("__DATA_JSON__", json.dumps(dataset, ensure_ascii=False, default=str))

    RUTA_HTML.write_text(html, encoding="utf-8")
    print(f"Dashboard generado en: {RUTA_HTML}")


if __name__ == "__main__":
    generar_html()
